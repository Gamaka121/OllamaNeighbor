from pathlib import Path

from openpyxl import Workbook
from docx import Document

from tools import files


DOCUMENTS_DIR = (
    Path.home() /
    "Documents" /
    "OllamaNeighbor"
)

DOCUMENTS_DIR.mkdir(
    parents=True,
    exist_ok=True
)


def _ok(message, **data):
    return {"success": True, "message": message, "data": data}


def _error(message, **data):
    return {"success": False, "error": message, "data": data}


def _document_path(filename, extension):
    raw = str(filename or "").strip()
    if not raw or raw in {".", ".."}:
        return None, "Укажите имя файла."
    if not raw.lower().endswith(extension):
        raw = raw + extension
    path, error = files._write_path(raw)
    if error:
        return None, error
    return path, None


def create_excel(filename, data):
    if not isinstance(data, (list, tuple)):
        return _error("Данные Excel должны быть списком строк.")

    if len(data) > 10_000:
        return _error("Слишком много строк для Excel (лимит 10 000).")

    path, error = _document_path(filename, ".xlsx")
    if error:
        return _error(error, path=str(filename or ""))

    workbook = Workbook()
    sheet = workbook.active

    for row in data:
        if not isinstance(row, (list, tuple)):
            return _error("Каждая строка Excel должна быть списком значений.")
        if len(row) > 200:
            return _error("В строке Excel слишком много ячеек (лимит 200).")
        sheet.append(list(row))

    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
    except OSError as error:
        return _error(f"Не удалось сохранить Excel-файл: {error}", path=str(path))

    return _ok("Excel-файл создан.", path=str(path), operation="create", type="file")


def create_word(filename, text):
    path, error = _document_path(filename, ".docx")
    if error:
        return _error(error, path=str(filename or ""))

    document = Document()

    for paragraph in str(text).split("\n"):
        document.add_paragraph(paragraph)

    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        document.save(path)
    except OSError as error:
        return _error(f"Не удалось сохранить Word-документ: {error}", path=str(path))

    return _ok("Word-документ создан.", path=str(path), operation="create", type="file")
